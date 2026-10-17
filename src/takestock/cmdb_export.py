"""Write a configuration register as JSON, CSV and a workbook.

Every writer takes the plain dict :func:`takestock.cmdb.build_register`
returns, so the three views are always of the same register:

* ``cmdb.json``: the whole register, schema first, then the data;
* ``csv/<class>.csv``: one file per declared class, every attribute the
  class carries as a column, header-only when the class is empty, plus
  ``csv/relationships.csv``;
* ``cmdb.xlsx``: a read-me, one sheet per populated class, the
  relationships, what is not in the register, disagreements, population
  counts, the per-fact provenance and the data dictionary;
* ``keys.jsonl``: the key ledger, for the next build to keep its keys;
* ``jira-assets/``, when asked for: an object schema and import mapping for
  Jira Assets, with the data to import (:func:`write_jira_assets`).

A missing value is written as its reason in brackets (``[not measured]``),
never as an empty cell. Styling is deliberately plain: bold headers on light
grey, absences in grey italics, contested values in one accent colour. Given
a :class:`takestock.htmlpage.Theme`, the workbook takes its colours and font
instead (:func:`workbook_style`).
"""

from __future__ import annotations

import csv
import json
import re
from collections.abc import Mapping
from dataclasses import dataclass
from pathlib import Path
from typing import Any

JSON_NAME = "cmdb.json"
XLSX_NAME = "cmdb.xlsx"
CSV_DIR = "csv"
KEYS_NAME = "keys.jsonl"
JIRA_DIR = "jira-assets"

_HEADER_FILL = "D9D9D9"
_ABSENT_INK = "767676"
_ACCENT = "C00000"

_ABSENT_WORDS = {
    "not-collected": "not collected",
    "not-measured": "not measured",
    "not-applicable": "not applicable",
    "could-not-test": "could not test",
    "extent-unknown": "extent unknown",
    "asked-unknown": "asked, unknown",
    "asked-declined": "asked, declined",
    "temp-unknown": "not known yet",
    "masked": "withheld",
    "unsupported": "source cannot hold it",
}


def absent_words(code: str) -> str:
    """An absence code as plain words."""
    return _ABSENT_WORDS.get(code, code.replace("-", " "))


def cell_text(cell: Mapping[str, Any] | None) -> Any:
    """What a spreadsheet or CSV shows for one attribute cell."""
    if cell is None:
        return ""
    if cell.get("absent"):
        return f"[{absent_words(str(cell['absent']))}]"
    value = cell.get("value")
    if isinstance(value, bool):
        return "yes" if value else "no"
    return value


def slug(name: str) -> str:
    """A class name as a file name: ``Software Release`` becomes ``software-release``."""
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def _compact(node: Any) -> Any:
    # Drop empty strings, nulls and false contest markers; keep every real
    # value, including 0 and False.
    if isinstance(node, dict):
        return {k: _compact(v) for k, v in node.items()
                if not (v == "" or v is None or (v is False and k in ("contested", "tie")))}
    if isinstance(node, list):
        return [_compact(v) for v in node]
    return node


def write_json(register: Mapping[str, Any], path: Path) -> Path:
    """The whole register as one JSON document, the schema beside the data."""
    body = {k: v for k, v in register.items() if k != "key_ledger"}
    path.write_text(json.dumps(_compact(body), indent=1, default=str, ensure_ascii=False) + "\n")
    return path


def write_keys(register: Mapping[str, Any], path: Path) -> Path:
    """The key ledger, one JSON record per line."""
    path.write_text("".join(json.dumps(r, sort_keys=True) + "\n"
                            for r in register.get("key_ledger") or []))
    return path


def _class_columns(register: Mapping[str, Any], name: str) -> list[str]:
    for spec in register["schema"]["classes"]:
        if spec["name"] == name:
            return [a for a in spec["attributes"] if a != "lifecycle_state"]
    return []


_ITEM_FIXED = ["key", "name", "lifecycle_state", "identifiers", "contested"]


def _item_row(item: Mapping[str, Any], columns: list[str]) -> list[Any]:
    cells = item["attributes"]
    contested = sorted(a for a, c in cells.items() if c.get("contested"))
    identifiers = "; ".join(f"{a['kind']}={a['identifier']}" for a in item["aliases"])
    return ([item["key"], item["name"], item["lifecycle_state"], identifiers,
             ", ".join(contested)] + [cell_text(cells.get(a)) for a in columns])


_RELATIONSHIP_COLUMNS = ["from", "from_name", "type", "to", "to_name", "count", "unit",
                         "evidence_strength", "valid_at", "qualifier"]


def write_csv(register: Mapping[str, Any], directory: Path) -> list[Path]:
    """One CSV per declared class, and one for the relationships."""
    directory.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for spec in register["schema"]["classes"]:
        columns = _class_columns(register, spec["name"])
        path = directory / f"{slug(spec['name'])}.csv"
        with path.open("w", newline="") as fh:
            out = csv.writer(fh)
            out.writerow(_ITEM_FIXED + columns)
            for item in register["items"]:
                if item["class"] == spec["name"]:
                    out.writerow(_item_row(item, columns))
        written.append(path)
    path = directory / "relationships.csv"
    with path.open("w", newline="") as fh:
        out = csv.writer(fh)
        out.writerow(_RELATIONSHIP_COLUMNS)
        for edge in register["relationships"]:
            out.writerow([edge.get(c, "") for c in _RELATIONSHIP_COLUMNS])
    written.append(path)
    return written


@dataclass(frozen=True)
class WorkbookStyle:
    """The colours (``RRGGBB``, or "" for the default ink) and font a workbook uses."""

    header_fill: str = _HEADER_FILL
    header_ink: str = ""
    heading_ink: str = ""
    absent_ink: str = _ABSENT_INK
    accent_ink: str = _ACCENT
    font: str = ""


# CSS generic and system font keywords, which name no font a workbook can use.
_GENERIC_FONTS = {"serif", "sans-serif", "monospace", "cursive", "fantasy", "system-ui",
                  "ui-serif", "ui-sans-serif", "ui-monospace", "ui-rounded", "math",
                  "emoji", "fangsong", "-apple-system", "blinkmacsystemfont"}


def _hex(colour: str) -> str:
    h = colour.lstrip("#")
    return "".join(c * 2 for c in h).upper() if len(h) == 3 else h.upper()


def workbook_style(theme: Any = None) -> WorkbookStyle:
    """The workbook's style: plain without a theme, else the theme's light colours and font.

    With a :class:`takestock.htmlpage.Theme`: headers in its ``text`` colour
    on its ``surface``, headings and contested values in its ``accent``,
    absences in its ``muted``, and the first named font in its font list
    (a generic family such as ``sans-serif`` names none, so the workbook
    keeps its default font).
    """
    if theme is None:
        return WorkbookStyle()
    light = theme.light
    font = ""
    for family in str(theme.font).split(","):
        name = family.strip().strip("'\"")
        if name and name.lower() not in _GENERIC_FONTS:
            font = name
            break
    return WorkbookStyle(header_fill=_hex(light["surface"]), header_ink=_hex(light["text"]),
                         heading_ink=_hex(light["accent"]), absent_ink=_hex(light["muted"]),
                         accent_ink=_hex(light["accent"]), font=font)


def write_xlsx(register: Mapping[str, Any], path: Path, *, theme: Any = None) -> Path:
    """The register as a workbook: plainly styled, or in ``theme``'s colours and font."""
    # openpyxl ships without type information.
    from openpyxl import Workbook  # type: ignore[import-untyped]
    from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
    from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

    style = workbook_style(theme)

    def font(**kw: Any) -> Any:
        if style.font:
            kw.setdefault("name", style.font)
        return Font(**kw)

    def ink(colour: str) -> dict[str, str]:
        return {"color": colour} if colour else {}

    wrap = Alignment(vertical="top", wrap_text=True)
    absent_font = font(color=style.absent_ink, italic=True)
    accent_font = font(color=style.accent_ink, bold=True)

    def table(ws: Any, names: list[str], rows: list[list[Any]], *, start: int = 1,
              widths: Mapping[str, int] | None = None) -> None:
        for col, name in enumerate(names, start=1):
            cell = ws.cell(row=start, column=col, value=name)
            cell.font = font(bold=True, **ink(style.header_ink))
            cell.fill = PatternFill("solid", fgColor=style.header_fill)
            cell.alignment = wrap
            ws.column_dimensions[get_column_letter(col)].width = (widths or {}).get(name, 20)
        for r, values in enumerate(rows, start=start + 1):
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=r, column=col, value=value)
                if isinstance(value, str) and value.startswith("[") and value.endswith("]"):
                    cell.font = absent_font
                elif style.font:
                    cell.font = font()
        ws.freeze_panes = f"A{start + 1}"
        ws.auto_filter.ref = f"A{start}:{get_column_letter(len(names))}{start + len(rows)}"

    reg = register["register"]
    counts = reg["counts"]
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Read me"
    lines: list[tuple[str, Any]] = [
        (str(reg["title"]), font(bold=True, size=16, **ink(style.heading_ink))),
    ]
    if reg.get("organisation"):
        lines.append((str(reg["organisation"]), font(bold=True)))
    lines += [
        (f"Observation {reg['observation']}, recorded {reg['recorded_at']}. "
         f"Schema {reg['schema_version']} ({reg['schema_maturity']}).", None),
        ("", None),
        (f"{counts['configuration_items']} configuration items, {counts['recorded_facts']} "
         f"recorded facts, {counts['relationships']} relationships. "
         f"{counts['classes_populated']} of {counts['classes_declared']} declared classes "
         f"hold items; the 'Not in here' sheet says why each other one is empty. "
         f"{counts['contested_attributes']} values are contested.", font(bold=True)),
    ]
    unassigned = reg["governance"].get("unassigned") or []
    if unassigned:
        lines.append(("No one is named for: " + ", ".join(unassigned) + ".", accent_font))
    lines.append(("", None))
    for entry in register["how_to_read"]:
        lines += [(entry["heading"], font(bold=True, **ink(style.heading_ink))),
                  (entry["body"], None), ("", None)]
    for r, (text, face) in enumerate(lines, start=1):
        cell = ws.cell(row=r, column=1, value=text or None)
        cell.alignment = wrap
        if face is not None:
            cell.font = face
        elif style.font:
            cell.font = font()
    ws.column_dimensions["A"].width = 110

    populated = [c["class"] for c in register["census"] if c["items"]]
    for name in populated:
        columns = _class_columns(register, name)
        sheet = wb.create_sheet(name[:31])
        members = [i for i in register["items"] if i["class"] == name]
        table(sheet, _ITEM_FIXED + columns, [_item_row(i, columns) for i in members],
              widths={"key": 12, "name": 34, "identifiers": 40, "contested": 18})
        for r, item in enumerate(members, start=2):
            for col, attribute in enumerate(columns, start=len(_ITEM_FIXED) + 1):
                if item["attributes"].get(attribute, {}).get("contested"):
                    sheet.cell(row=r, column=col).font = accent_font

    table(wb.create_sheet("Relationships"), _RELATIONSHIP_COLUMNS,
          [[e.get(c, "") for c in _RELATIONSHIP_COLUMNS] for e in register["relationships"]],
          widths={"from_name": 34, "to_name": 34, "unit": 30, "qualifier": 60})

    gaps = wb.create_sheet("Not in here")
    table(gaps, ["class", "family", "items", "answers", "built_from", "empty_because"],
          [[c["class"], c["family"], c["items"], c["answers"], c["source"], c["empty_because"]]
           for c in register["census"]],
          widths={"class": 24, "answers": 50, "built_from": 40, "empty_because": 70})
    start = len(register["census"]) + 3
    table(gaps, ["relationship not drawn", "why"],
          [[r["relationship"], r["reason"]] for r in register["relationships_not_drawn"]],
          start=start, widths={"relationship not drawn": 24, "why": 50})
    gaps.freeze_panes = "A2"
    gaps.auto_filter.ref = f"A1:F{len(register['census']) + 1}"

    vocab = register["schema"]["vocabularies"].get("absence_reason", {})
    recorded = register["absence"]["recorded_by_reason"]
    table(wb.create_sheet("Absence counted"), ["absence", "values", "meaning"],
          [[absent_words(code), n, str(vocab.get(code, {}).get("definition", ""))]
           for code, n in sorted(recorded.items())]
          + [["declared and unasserted", register["absence"]["attributes_declared_and_unasserted"],
              "Attributes a class declares that no row stated. Counted apart from the rows "
              "above and never added to them."]],
          widths={"absence": 26, "values": 10, "meaning": 90})

    def said(value: Any) -> str:
        if isinstance(value, Mapping) and "absent" in value:
            return f"[{absent_words(str(value['absent']))}]"
        return "" if value is None else str(value)

    table(wb.create_sheet("Disagreements"),
          ["key", "attribute", "published value", "evidence_strength", "valid_at",
           "asserted_by", "tie", "what each source said"],
          [[c["item"], c["attribute"], said(c["value"]), c["evidence_strength"], c["valid_at"],
            c["actor"], "yes" if c["tie"] else "",
            " | ".join(f"{said(a['value'])} ({a['evidence_strength']}, {a['valid_at']}, "
                       f"{a['actor']})" for a in c["assertions"])]
           for c in register["conflicts"]],
          widths={"published value": 30, "asserted_by": 30, "what each source said": 90})

    table(wb.create_sheet("Populations"),
          ["population", "instrument", "count", "unit", "valid_at", "evidence_strength", "note"],
          [[p["population"], p["instrument"], p["count"], p["unit"], p["valid_at"],
            p["evidence_strength"], p["note"]] for p in register["populations"]],
          widths={"instrument": 40, "note": 70})

    fact_columns = ["key", "class", "name", "attribute", "value", "evidence_strength",
                    "valid_at", "recorded_at", "asserted_by", "source_artefact", "basis",
                    "contested"]
    table(wb.create_sheet("Facts"), fact_columns,
          [[f["key"], f["class"], f["name"], f["attribute"],
            f"[{absent_words(f['absent'])}]" if f["absent"] else cell_text(f),
            f["evidence_strength"], f["valid_at"], f["recorded_at"], f["asserted_by"],
            f["source_artefact"], f["basis"], "yes" if f["contested"] else ""]
           for f in register["facts"]],
          widths={"name": 30, "value": 34, "asserted_by": 30, "source_artefact": 30})

    dictionary = wb.create_sheet("Data dictionary")
    table(dictionary, ["attribute", "expected_strength", "refresh_interval_days",
                       "carried_by", "note"],
          [[a["name"], a["expected_strength"], a["refresh_interval_days"],
            ", ".join(a["carried_by"]), a["note"]] for a in register["schema"]["attributes"]],
          widths={"attribute": 26, "carried_by": 50, "note": 90})

    wb.save(path)
    return path


# ── An import for Jira Assets ──────────────────────────────────────────────
#
# Jira Assets (Atlassian's asset and configuration database) takes an external
# import as two JSON documents: an object schema with its import mapping, and
# the data. Their shape is Atlassian's published "Imports Schema and Mapping
# definition", version 2023-10-19, which Assets applies to any document that
# uses external ids, as every one written here does; the data is submitted as
# one chunk whose "data" object holds a list of records per mapping selector.

#: The version of Atlassian's import schema-and-mapping definition written.
JIRA_SCHEMA_VERSION = "2023-10-19"
JIRA_MAPPING_NAME = "schema-and-mapping.json"
JIRA_DATA_NAME = "data.json"
#: Every Assets object has a key of its own, so the register's key takes this name.
JIRA_KEY_ATTRIBUTE = "Register key"

_JIRA_FIXED: tuple[tuple[str, str, str, str], ...] = (
    # (locator, attribute name, Assets type, description)
    ("register_key", JIRA_KEY_ATTRIBUTE, "text",
     "The register's surrogate key: opaque, never reused, and the attribute an import "
     "matches on. It names the same item from one import to the next only when each "
     "build is given the previous build's keys.jsonl."),
    ("register_name", "Name", "text", "The item's name in the register."),
    ("register_lifecycle_state", "Lifecycle state", "text",
     "Set by the build: discovered when an instrument saw it, planned when it exists "
     "only on paper."),
    ("register_identifiers", "Identifiers", "textarea",
     "Every identifier the item was seen under, written kind=identifier."),
    ("register_contested", "Contested", "text",
     "The attributes whose sources disagree; the register's Disagreements sheet says what "
     "each source said."),
)


_JIRA_LOCATORS = tuple(f[0] for f in _JIRA_FIXED)
_JIRA_LABELS = {f[1].lower() for f in _JIRA_FIXED}


def _jira_part(name: str) -> str:
    return slug(name).replace("-", "_")


def _jira_text(value: Any) -> str:
    if isinstance(value, list):
        return "; ".join(_jira_text(v) for v in value)
    if isinstance(value, Mapping):
        return json.dumps(value, sort_keys=True, default=str, ensure_ascii=False)
    return "" if value is None else str(value)


def jira_assets(register: Mapping[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    """The register as a Jira Assets external import: (schema and mapping, data chunk).

    One object type per declared class. Each carries the register key (the
    attribute objects are matched on, ``externalIdPart``), the name (the
    object's label), the lifecycle state, the identifiers, which attributes
    are contested, and every attribute the class declares, as text, absences
    written as their reason in brackets. Each kind of drawn relationship
    between two classes becomes a reference attribute on the class it starts
    from, holding every key it points at and resolved with an AQL ``IN``
    query on the register key.
    """
    reg = register["register"]
    classes = register["schema"]["classes"]
    by_key = {i["key"]: i for i in register["items"]}
    references: dict[str, dict[tuple[str, str], str]] = {}
    targets: dict[str, dict[str, list[str]]] = {}
    for edge in register["relationships"]:
        source, target = by_key.get(edge["from"]), by_key.get(edge["to"])
        if source is None or target is None:
            continue
        locator = f"{_jira_part(edge['type'])}__{_jira_part(target['class'])}"
        references.setdefault(source["class"], {})[(edge["type"], target["class"])] = locator
        pointed = targets.setdefault(source["key"], {}).setdefault(locator, [])
        if target["key"] not in pointed:
            pointed.append(target["key"])

    object_types: list[dict[str, Any]] = []
    mappings: list[dict[str, Any]] = []
    data: dict[str, list[dict[str, Any]]] = {}
    for spec in classes:
        name = spec["name"]
        part = _jira_part(name)
        type_id = f"takestock/{part}"
        columns = [a for a in spec["attributes"] if a != "lifecycle_state"]
        notes = {a["name"]: a["note"] for a in register["schema"]["attributes"]}
        attributes: list[dict[str, Any]] = []
        mapped: list[dict[str, Any]] = []
        for locator, label, kind, description in _JIRA_FIXED:
            attribute: dict[str, Any] = {"externalId": f"{type_id}/{locator}", "name": label,
                                         "description": description, "type": kind}
            if locator == "register_key":
                attribute["unique"] = True
            if locator == "register_name":
                attribute["label"] = True
            attributes.append(attribute)
            mapped.append({"attributeExternalId": f"{type_id}/{locator}", "attributeName": label,
                           "attributeLocators": [locator],
                           "externalIdPart": locator == "register_key"})
        for column in columns:
            if column in _JIRA_LOCATORS:
                raise ValueError(f"attribute {column!r} of {name} has the name of a fixed "
                                 "Jira Assets column")
            # Assets attribute names are compared without regard to case.
            label = f"{column} (attribute)" if column.lower() in _JIRA_LABELS else column
            attributes.append({"externalId": f"{type_id}/{column}", "name": label,
                               "description": notes.get(column) or column, "type": "text"})
            mapped.append({"attributeExternalId": f"{type_id}/{column}", "attributeName": label,
                           "attributeLocators": [column], "externalIdPart": False})
        for (kind_of, to_class), locator in sorted(references.get(name, {}).items()):
            label = f"{kind_of} {to_class}"
            attributes.append({"externalId": f"{type_id}/{locator}", "name": label,
                               "description": f"The {to_class} items this one is {kind_of}.",
                               "type": "referenced_object",
                               "referenceObjectTypeExternalId": f"takestock/{_jira_part(to_class)}",
                               "minimumCardinality": 0, "maximumCardinality": -1})
            mapped.append({"attributeExternalId": f"{type_id}/{locator}", "attributeName": label,
                           "attributeLocators": [locator], "externalIdPart": False,
                           "objectMappingIQL": f'"{JIRA_KEY_ATTRIBUTE}" IN (${{{locator}${{0}}}})'})
        description = spec["answers"] or f"{name} items in the configuration register."
        object_types.append({"externalId": type_id, "name": name, "description": description,
                             "attributes": attributes})
        mappings.append({"objectTypeExternalId": type_id, "objectTypeName": name,
                         "selector": part, "description": f"{name} items, one record each.",
                         "attributesMapping": mapped})
        records: list[dict[str, Any]] = []
        for item in register["items"]:
            if item["class"] != name:
                continue
            row = _item_row(item, columns)
            record: dict[str, Any] = dict(zip(_JIRA_LOCATORS, row[:5], strict=True))
            record.update({c: _jira_text(v) for c, v in zip(columns, row[5:], strict=True)})
            record.update(targets.get(item["key"], {}))
            records.append(record)
        data[part] = records

    who = f" for {reg['organisation']}" if reg.get("organisation") else ""
    mapping = {
        "schema": {"objectSchema": {
            "name": str(reg["title"]),
            "description": (f"Draft configuration register{who}, built by takestock from "
                            f"observation {reg['observation']}."),
            "objectTypes": object_types}},
        "mapping": {"objectTypeMappings": mappings},
    }
    return mapping, {"data": data, "completed": True}


def write_jira_assets(register: Mapping[str, Any], directory: Path) -> list[Path]:
    """Write :func:`jira_assets`'s two documents into ``directory``; return their paths."""
    directory.mkdir(parents=True, exist_ok=True)
    mapping, chunk = jira_assets(register)
    out = []
    for name, body in ((JIRA_MAPPING_NAME, mapping), (JIRA_DATA_NAME, chunk)):
        path = directory / name
        path.write_text(json.dumps(body, indent=1, ensure_ascii=False) + "\n")
        out.append(path)
    return out


def write_all(register: Mapping[str, Any], out_dir: Path, *, theme: Any = None,
              jira_assets: bool = False) -> list[Path]:
    """Write every view of ``register`` into ``out_dir``; return the paths written.

    ``theme`` styles the workbook (:func:`workbook_style`); ``jira_assets``
    also writes the Jira Assets import into ``jira-assets/``.
    """
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = ([write_json(register, out_dir / JSON_NAME)]
             + write_csv(register, out_dir / CSV_DIR)
             + [write_xlsx(register, out_dir / XLSX_NAME, theme=theme),
                write_keys(register, out_dir / KEYS_NAME)])
    if jira_assets:
        paths += write_jira_assets(register, out_dir / JIRA_DIR)
    return paths
