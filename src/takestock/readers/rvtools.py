"""RVTools workbooks: the shape of a VMware virtual estate, with no encryption in it.

**RVTools** is a free Windows tool that reads a vCenter or ESXi host through
the vSphere API and writes what it finds to a workbook, one sheet per
subject. Its default file name is
``RVTools_export_all_<YYYY-MM-DD>_<HH.MM.SS>.xlsx``; with more than one
vCenter, the person exporting usually puts a label before it
(``SiteA_RVTools_export_all_…``). That label is the row's **site**, and the
date in the name is the export's date (the file's modification date when the
name has none).

**The input** is one workbook or a folder of them (every ``.xlsx`` file
directly inside it, in name order). Two sheets are read:

* ``vInfo``, one row per virtual machine: ``VM``, ``Powerstate``,
  ``Template``, ``SRM Placeholder``, ``Firmware``, ``EFI Secure boot`` and
  ``OS according to the configuration file``;
* ``vHost``, one row per host: ``Host`` and ``ESX Version``.

**There is no encryption data in an RVTools workbook**: no sheet has a column
for VM encryption, a key provider or a cipher. Every row here is a census,
graded ``n/a``, and the scope row says so, so an inventory is never read as
evidence about data at rest.

**A row is not always a machine.** Templates, and Site Recovery Manager
placeholders (a stand-in registered at the recovery site for a machine that
runs at the other), are rows in ``vInfo`` and are not counted as machines.

**Rows.**

* Over the machines powered on, one per **firmware and secure-boot** state:
  EFI with secure boot on; EFI with it off (``secure-boot-off:<n>``); legacy
  BIOS (``legacy-bios:<n>``), counted apart because it cannot do secure boot
  at all, so turning it on needs a firmware change, not a setting; and any
  other firmware value. Each carries ``vm-firmware-census:<n>``.
* One per **ESXi version** on the hosts (``esxi-version-census:<n>``).
* One per configured **guest operating system** (``guest-os-census:<n>``),
  as set in each machine's settings, not as measured inside it.
* One **scope** row, ``Virtual estate scope``
  (``scope-rvtools-no-encryption-data``), with the reconciliation from rows
  to machines.

**With the vSphere Client's list exports as well.** Those lists
(``vsphere_lists``) count machines, hosts, guest types and ESXi versions
straight from vCenter. When this reader's ``settings`` name that reader's
folder under ``vsphere_lists`` and it holds a virtual-machine list, only the
firmware and secure-boot rows are written here, the one thing the lists do
not carry, with a shorter scope row saying so.

**Dates.** Every row is dated by the oldest workbook's date and its notes say
it describes the estate then, because a workbook is a snapshot.

**Cannot see.** Encryption of any kind; what runs inside a guest; and
anything about a vCenter not exported.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any

from .. import coverage, roles
from ..model import CryptoAsset, register_informational_prefixes
from ..plugins import Context, Plugin, StepResult
from ..util import file_mtime_date
from .vsphere_lists import lists_folder, vm_lists_present

#: The plug-in's name, and so the ``source`` of every row it writes.
NAME = "rvtools"

CENSUS_FLAG = "vm-firmware-census:"
SECURE_BOOT_OFF_FLAG = "secure-boot-off:"
LEGACY_BIOS_FLAG = "legacy-bios:"
ESXI_FLAG = "esxi-version-census:"
GUEST_OS_FLAG = "guest-os-census:"
SCOPE_FLAG = "scope-rvtools-no-encryption-data"

register_informational_prefixes(CENSUS_FLAG, SECURE_BOOT_OFF_FLAG, LEGACY_BIOS_FLAG, ESXI_FLAG,
                                GUEST_OS_FLAG, SCOPE_FLAG)

#: The ``vInfo`` column holding the guest type a machine is configured as.
GUEST_OS_COLUMN = "OS according to the configuration file"

_DATE_IN_NAME = re.compile(r"(\d{4}-\d{2}-\d{2})")


def site_of(path: Path) -> str:
    """The label before ``_RVTools`` in the file name, or the whole name when there is none."""
    stem = path.stem
    return stem.split("_RVTools")[0] if "_RVTools" in stem else stem


def date_of(path: Path) -> str:
    """The date RVTools writes into the file name, else the file's modification date."""
    m = _DATE_IN_NAME.search(path.name)
    return m.group(1) if m else (file_mtime_date(path) or "")


def _true(value: object) -> bool:
    return str(value).strip().lower() == "true"


def _on(row: dict[str, Any]) -> bool:
    return str(row.get("Powerstate")) == "poweredOn"


def sheet(workbook: Any, name: str) -> list[dict[str, Any]]:
    """A sheet's rows keyed by its first row's headings; [] when the sheet is absent."""
    if name not in workbook.sheetnames:
        return []
    cells = workbook[name].iter_rows(values_only=True)
    head = [str(h) if h is not None else "" for h in next(cells, [])]
    return [{h: raw[i] if i < len(raw) else None for i, h in enumerate(head)} for raw in cells]


def workbooks(path: Path) -> list[Path]:
    """The workbook named, or every ``.xlsx`` directly inside a folder, in name order."""
    if path.is_dir():
        return sorted(p for p in path.glob("*.xlsx") if p.is_file())
    return [path]


def read(ctx: Context) -> StepResult:
    """The plug-in's step.

    With no input, nothing is written: an inventory of the virtual estate
    answers no cryptographic question by itself, so its absence is not a
    gap. A workbook that cannot be read is.
    """
    if ctx.input is None:
        return StepResult()
    # openpyxl ships without type information.
    import openpyxl  # type: ignore[import-untyped]
    from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]

    sites: list[dict[str, Any]] = []
    unreadable: list[str] = []
    found = workbooks(ctx.input)
    if not found:
        unreadable.append(f"{ctx.input.name} (no .xlsx workbook in the folder)")
    for path in found:
        try:
            wb = openpyxl.load_workbook(path, read_only=True)
        except (OSError, KeyError, ValueError, zipfile.BadZipFile, InvalidFileException) as exc:
            unreadable.append(f"{path.name} ({exc})")
            continue
        try:
            vinfo, vhost = sheet(wb, "vInfo"), sheet(wb, "vHost")
        finally:
            wb.close()
        if not vinfo:
            unreadable.append(f"{path.name} (no vInfo rows, so not an RVTools export)")
            continue
        sites.append({"site": site_of(path), "date": date_of(path), "vinfo": vinfo,
                      "vhost": vhost})
    gaps: list[coverage.Gap] = []
    if unreadable:
        gaps.append(coverage.gap(
            coverage.NOT_COLLECTED,
            "RVTools workbooks that could not be read: " + "; ".join(unreadable),
            impact_kind=coverage.NARROWS_CONFIDENCE,
            impact=("These rows cover only the workbooks that could be read, so every count "
                    "may be short of the whole virtual estate."),
            harvest="Export the workbook from vCenter again with RVTools and replace the file.",
            nature=coverage.NO_DATA))
    if not sites:
        return StepResult(gaps=gaps)
    superseded = vm_lists_present(lists_folder(ctx))
    return StepResult(records=_rows(sites, superseded), gaps=gaps)


def _asset(**kw: Any) -> CryptoAsset:
    return CryptoAsset(data_state="platform-integrity", protocol="firmware", source=NAME,
                       source_class="private", verdict="n/a", **kw)


def _rows(sites: list[dict[str, Any]], superseded: bool) -> list[CryptoAsset]:
    dates = sorted({s["date"] for s in sites if s["date"]})
    stamp = dates[0] if dates else ""
    dates_text = " and ".join(dates) if dates else "an unstated date"
    names = [s["site"] for s in sites]
    source_name = f"RVTools workbooks {', '.join(names)}, dated {dates_text}"
    snapshot = (f"From RVTools workbooks dated {dates_text}: the virtual estate as it was then, "
                "not as it is now.")

    counts = {"rows": 0, "templates": 0, "placeholders": 0}
    machines: list[dict[str, Any]] = []
    per_site: dict[str, dict[str, int]] = {}
    for s in sites:
        t = per_site.setdefault(s["site"], {"rows": 0, "templates": 0, "placeholders": 0,
                                            "machines": 0, "on": 0})
        for r in s["vinfo"]:
            if r.get("VM") in (None, ""):
                continue
            counts["rows"] += 1
            t["rows"] += 1
            if _true(r.get("Template")):
                counts["templates"] += 1
                t["templates"] += 1
                continue
            if _true(r.get("SRM Placeholder")):
                counts["placeholders"] += 1
                t["placeholders"] += 1
                continue
            t["machines"] += 1
            t["on"] += _on(r)
            machines.append({**r, "_site": s["site"]})
    running = [r for r in machines if _on(r)]

    out: list[CryptoAsset] = []
    by_firmware: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for r in running:
        fw = str(r.get("Firmware") or "").strip().lower() or "(not reported)"
        key = (fw, ("on" if _true(r.get("EFI Secure boot")) else "off") if fw == "efi" else "")
        by_firmware.setdefault(key, []).append(r)
    for (fw, sb), members in sorted(by_firmware.items(), key=lambda kv: -len(kv[1])):
        n = len(members)
        of = f"{n} of the {len(running)} machines powered on"
        flags = [f"{CENSUS_FLAG}{n}"]
        if fw == "efi" and sb == "off":
            name = "VM firmware EFI, secure boot off"
            flags.append(f"{SECURE_BOOT_OFF_FLAG}{n}")
            what = (f"{of} have EFI firmware with secure boot off: they can verify what they "
                    "start, and the setting is not on. ")
        elif fw == "efi":
            name = "VM firmware EFI, secure boot on"
            what = (f"{of} have EFI firmware with secure boot on: the firmware checks what "
                    "they load before the operating system starts. ")
        elif fw == "bios":
            name = "VM firmware BIOS (legacy)"
            flags.append(f"{LEGACY_BIOS_FLAG}{n}")
            what = (f"{of} have legacy BIOS firmware, which cannot do secure boot: turning "
                    "it on needs a change of firmware, not of a setting. ")
        else:
            name = f"VM firmware {fw}"
            what = f"{of} report their firmware as {fw!r}. "
        out.append(_asset(
            name=name, role=roles.virtual_machine_boot_population(n), source_name=source_name,
            collected_at=stamp, flags=flags, hosts_count=n, notes=what + snapshot,
            extra={"firmware": fw, "secure_boot": sb or "n/a",
                   "powered_on_total": str(len(running)),
                   "sites": ", ".join(sorted({m["_site"] for m in members}))}))

    if superseded:
        out.append(_asset(
            name="Virtual estate scope", role=roles.virtual_estate_scope(),
            source_name=source_name, collected_at=stamp, flags=[SCOPE_FLAG],
            notes=(f"The firmware rows count the {len(running)} machines powered on in the "
                   f"workbooks dated {dates_text}, with templates and Site Recovery Manager "
                   "placeholders left out. The vSphere Client's list exports are read for "
                   "the counts of machines, hosts, guest types and ESXi versions, so none is "
                   "written from these workbooks; firmware and secure boot are what the lists "
                   "do not carry. An RVTools workbook has no encryption column, and VM "
                   "encryption is read from the lists. " + snapshot),
            extra={"superseded_by": "vsphere_lists", "powered_on": str(len(running)),
                   "export_dates": dates_text, "sites": ", ".join(names)}))
        return out

    versions: dict[str, list[tuple[str, str]]] = {}
    for s in sites:
        for r in s["vhost"]:
            host = str(r.get("Host") or "").strip()
            if host:
                version = str(r.get("ESX Version") or "").strip() or "(not reported)"
                versions.setdefault(version, []).append((s["site"], host))
    hosts = sum(len(v) for v in versions.values())
    for version, pairs in sorted(versions.items(), key=lambda kv: -len(kv[1])):
        n = len(pairs)
        out.append(_asset(
            name=version, role=roles.hypervisor_host_population(n), software=version,
            source_name=source_name, collected_at=stamp, flags=[f"{ESXI_FLAG}{n}"], hosts_count=n,
            notes=(f"{n} of the {hosts} hosts ran {version} on the export date. No support "
                   "status is given here: that needs a published lifecycle, which this "
                   "inventory is not. " + snapshot),
            extra={"esx_version": version, "hosts": " ".join(sorted(h for _s, h in pairs)),
                   "sites": ", ".join(sorted({s for s, _h in pairs}))}))

    guests: dict[str, dict[str, int]] = {}
    for r in machines:
        g = guests.setdefault(str(r.get(GUEST_OS_COLUMN) or "").strip() or "(not recorded)",
                              {"vms": 0, "on": 0})
        g["vms"] += 1
        g["on"] += _on(r)
    for os_name, g in sorted(guests.items(), key=lambda kv: -kv[1]["vms"]):
        n = g["vms"]
        out.append(_asset(
            name=f"Guest OS {os_name}", role=roles.virtual_guest_os_population(n),
            source_name=source_name, collected_at=stamp, flags=[f"{GUEST_OS_FLAG}{n}"],
            hosts_count=n,
            notes=(f"{n} of the {len(machines)} machines were configured as this operating "
                   f"system on the export date, {g['on']} of them powered on. This is the "
                   "type set in each machine's settings, not measured inside the guest. "
                   + snapshot),
            extra={"guest_os": os_name, "vms": str(n), "powered_on": str(g["on"])}))

    windows = sum(1 for r in machines
                  if str(r.get(GUEST_OS_COLUMN) or "").lower().startswith("microsoft windows"))
    distinct = len({str(r.get("VM")) for r in machines})
    by_site = "; ".join(
        f"{name}: {t['rows']} rows less {t['templates']} template(s) and {t['placeholders']} "
        f"placeholder(s) is {t['machines']} machines, {t['on']} powered on"
        for name, t in per_site.items())
    out.append(_asset(
        name="Virtual estate scope", role=roles.virtual_estate_scope(),
        source_name=source_name, collected_at=stamp, flags=[SCOPE_FLAG],
        hosts_count=len(machines),
        notes=(f"{counts['rows']} vInfo rows in {len(sites)} workbook(s) dated {dates_text}. "
               f"{counts['templates']} templates and {counts['placeholders']} Site Recovery "
               f"Manager placeholders are rows, not machines, leaving {len(machines)} "
               f"machines ({distinct} distinct names), {len(running)} powered on, on {hosts} "
               f"hosts. By site: {by_site}. {windows} of the {len(machines)} are configured as "
               "Microsoft Windows. An RVTools workbook has no column for VM encryption, a key "
               "provider or a cipher, so nothing here is evidence about data at rest. "
               + snapshot),
        extra={"vinfo_rows": str(counts["rows"]), "templates": str(counts["templates"]),
               "srm_placeholders": str(counts["placeholders"]), "real_vms": str(len(machines)),
               "distinct_vm_names": str(distinct), "powered_on": str(len(running)),
               "windows": str(windows), "non_windows": str(len(machines) - windows),
               "esxi_hosts": str(hosts), "export_dates": dates_text, "sites": ", ".join(names)}))
    return out


#: The declaration. It names no default input; give a workbook or a folder of
#: them with ``input:``.
PLUGIN = Plugin(
    name=NAME, kind="general", step=read,
    title="RVTools workbooks: virtual machines' firmware, guest types and ESXi versions",
    gap_prefix=False)
