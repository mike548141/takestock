"""Requirements derived from an observation's findings.

``data/requirements-catalogue.yaml`` holds generic requirement templates, each
written from public standards and each carrying a ``triggered_by`` rule over
findings. :func:`derive` joins the catalogue to an observation's graded assets
(and their dispositions, where set) and to its coverage gaps, and returns only
the requirements at least one finding triggers. Each one carries an evidence
sentence written from the observation's own counts and hosts.

:func:`write_outputs` writes the result three ways: a markdown register, a
spreadsheet register with empty supplier-response columns, and JSON.

Run ``python -m takestock.requirements <observation_dir> <out_dir>``; the
observation directory holds ``assets.json`` and, optionally, ``summary.json``
with the coverage gaps.
"""

from __future__ import annotations

import json
import sys
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from . import config, coverage
from .exits import guarded
from .model import CryptoAsset, load_assets
from .policy import Policy

CATALOGUE_FILE = "requirements-catalogue.yaml"

LEVELS = ("MUST", "SHOULD", "COULD")

#: Keys an asset clause may use. All keys in one clause must hold for the same asset.
ASSET_KEYS = frozenset({
    "verdict", "data_state", "protocol", "source", "disposition", "internet_facing",
    "flags_any", "flags_all", "flags_none", "flags_prefix_any", "flags_prefix_none", "extra_true",
    "versions_none",
})
_CLAUSE_KINDS = ("absent", "gap_kinds")
_REQUIRED_FIELDS = ("id", "level", "area", "name", "statement", "standards",
                    "selection_test", "acceptance_test", "triggered_by")

#: How many hosts or gap examples an evidence sentence names before "and N more".
NAMED_LIMIT = 5

# Supplier-response columns left empty in the register for a bidder to fill in.
RESPONSE_COLUMNS = ("Supplier response (comply / partly / do not comply)",
                    "Supplier explanation", "Supplier evidence reference")
RESPONSE_CHOICES = "Comply,Partly comply,Do not comply"


@dataclass
class Match:
    """What one trigger clause matched: its finding phrase and the rows behind it."""

    finding: str
    unit: str
    count: int
    hosts: list[str] = field(default_factory=list)
    refs: list[str] = field(default_factory=list)
    examples: list[str] = field(default_factory=list)


@dataclass
class Requirement:
    """A triggered requirement: the catalogue template plus the evidence for it."""

    id: str
    level: str
    area: str
    name: str
    statement: str
    standards: list[dict]
    selection_test: str
    acceptance_test: str
    matches: list[Match]

    @property
    def evidence(self) -> str:
        """One sentence naming each matched clause with its count and hosts."""
        return evidence_sentence(self.matches)

    @property
    def findings_count(self) -> int:
        """Rows behind the requirement, summed over its clauses."""
        return sum(m.count for m in self.matches)

    def to_dict(self) -> dict:
        """The requirement as plain data, for JSON."""
        return {
            "id": self.id, "level": self.level, "area": self.area, "name": self.name,
            "statement": self.statement, "standards": self.standards,
            "selection_test": self.selection_test, "acceptance_test": self.acceptance_test,
            "evidence": self.evidence,
            "findings": [{"finding": m.finding, "unit": m.unit, "count": m.count,
                          "hosts": m.hosts, "refs": m.refs, "examples": m.examples}
                         for m in self.matches],
        }


# ── catalogue ──

def load_catalogue(path: str | Path | None = None) -> list[dict]:
    """Read and validate the catalogue; return its requirement templates in order.

    Defaults to the workspace's ``data/`` copy, else the bundled one. Raises
    ValueError naming the entry for anything a register could not be built
    from honestly: a missing field, an unknown level, a duplicate id, a
    trigger with no finding phrase, or a clause key nothing reads.
    """
    p = Path(path) if path else config.data_file(CATALOGUE_FILE)
    doc = config.load_yaml_strict(p)
    return validate_catalogue(doc.get("requirements") or [])


def validate_catalogue(entries: Sequence[Mapping[str, Any]]) -> list[dict]:
    """Check catalogue entries and return them as plain dicts (see :func:`load_catalogue`)."""
    seen: set[str] = set()
    out = []
    for n, e in enumerate(entries):
        rid = str(e.get("id") or f"#{n}")
        missing = [k for k in _REQUIRED_FIELDS if not e.get(k)]
        if missing:
            raise ValueError(f"requirement {rid}: missing {missing}")
        if rid in seen:
            raise ValueError(f"requirement {rid}: duplicate id")
        seen.add(rid)
        if e["level"] not in LEVELS:
            raise ValueError(f"requirement {rid}: level {e['level']!r} is not one of {LEVELS}")
        for s in e["standards"]:
            if not all(s.get(k) for k in ("ref", "level", "says")):
                raise ValueError(f"requirement {rid}: a standard needs ref, level and says: {s}")
        for clause in e["triggered_by"]:
            _check_clause(rid, clause)
        out.append(dict(e))
    return out


def _check_clause(rid: str, clause: Mapping[str, Any]) -> None:
    if not str(clause.get("finding") or "").strip():
        raise ValueError(f"requirement {rid}: a trigger clause has no finding phrase: {clause}")
    keys = set(clause) - {"finding"}
    if not keys:
        raise ValueError(f"requirement {rid}: a trigger clause matches nothing: {clause}")
    kinds = keys & set(_CLAUSE_KINDS)
    if kinds and keys != kinds or len(kinds) > 1:
        raise ValueError(f"requirement {rid}: {sorted(kinds)} cannot be combined with other "
                         f"keys in one clause: {clause}")
    if "absent" in kinds:
        a = clause["absent"] or {}
        if not (a.get("anchor") and a.get("missing")):
            raise ValueError(f"requirement {rid}: absent needs anchor and missing: {clause}")
    if "gap_kinds" in kinds:
        bad = [k for k in clause["gap_kinds"] if k not in coverage.KINDS]
        if bad:
            raise ValueError(f"requirement {rid}: unknown gap kinds {bad}")
    unknown = keys - ASSET_KEYS - set(_CLAUSE_KINDS)
    if unknown:
        raise ValueError(f"requirement {rid}: unknown trigger keys {sorted(unknown)}")


# ── matching ──

def _as_list(v: Any) -> list:
    return list(v) if isinstance(v, (list, tuple, set)) else [v]


def asset_matches(a: CryptoAsset, clause: Mapping[str, Any]) -> bool:
    """True when every asset key in ``clause`` holds for ``a``."""
    flags = set(a.flags)
    for key in ("verdict", "data_state", "protocol", "source", "disposition"):
        if key in clause and getattr(a, key) not in _as_list(clause[key]):
            return False
    if "internet_facing" in clause and a.internet_facing is not clause["internet_facing"]:
        return False
    if "flags_any" in clause and not flags & set(clause["flags_any"]):
        return False
    if "flags_all" in clause and not set(clause["flags_all"]) <= flags:
        return False
    if "flags_none" in clause and flags & set(clause["flags_none"]):
        return False
    if "flags_prefix_any" in clause:
        prefixes = tuple(clause["flags_prefix_any"])
        if not any(f.startswith(prefixes) for f in flags):
            return False
    if "flags_prefix_none" in clause:
        prefixes = tuple(clause["flags_prefix_none"])
        if any(f.startswith(prefixes) for f in flags):
            return False
    if "extra_true" in clause and not (a.extra or {}).get(clause["extra_true"]):
        return False
    if "versions_none" in clause:
        offered = {Policy.normalise_tls_version(v) for v in a.protocol_versions}
        barred = {Policy.normalise_tls_version(v) for v in clause["versions_none"]}
        if not offered or offered & barred:
            return False
    return True


def _host_label(a: CryptoAsset) -> str:
    host = a.host or a.ip or a.name
    return f"{host}:{a.port}" if a.port else host


def _norm_host(h: str) -> str:
    return h.lower().strip().rstrip(".")


def _match_clause(clause: Mapping[str, Any], assets: Sequence[CryptoAsset],
                  gaps: Sequence[object]) -> Match | None:
    finding = " ".join(str(clause["finding"]).split())
    if "gap_kinds" in clause:
        kinds = set(clause["gap_kinds"])
        hit = [g for g in gaps if coverage.kind_of(g) in kinds]
        if not hit:
            return None
        return Match(finding, "coverage gap", len(hit),
                     examples=[_clip(str(g)) for g in hit])
    if "absent" in clause:
        anchor, missing = clause["absent"]["anchor"], clause["absent"]["missing"]
        # A row recording a measured absence (extra.present False) is not the record.
        have = {_norm_host(a.host) for a in assets
                if a.protocol == missing and a.host and a.extra.get("present") is not False}
        lacking = sorted({_norm_host(a.host) for a in assets
                          if a.protocol == anchor and a.host} - have)
        if not lacking:
            return None
        return Match(finding, "domain", len(lacking), hosts=lacking)
    hit_assets = [a for a in assets if asset_matches(a, clause)]
    if not hit_assets:
        return None
    return Match(finding, "asset", len(hit_assets),
                 hosts=sorted({_host_label(a) for a in hit_assets}),
                 refs=[a.bom_ref() for a in hit_assets])


def _clip(text: str, limit: int = 100) -> str:
    flat = " ".join(text.split())
    return flat if len(flat) <= limit else flat[: limit - 1].rstrip() + "…"


def derive(assets: Iterable[CryptoAsset], catalogue: Sequence[Mapping[str, Any]],
           gaps: Iterable[object] = ()) -> list[Requirement]:
    """The catalogue requirements the findings trigger, in catalogue order.

    ``gaps`` are the observation's coverage gaps (:class:`coverage.Gap` or
    plain strings, which read as unclassified). A requirement none of whose
    clauses matches is left out entirely.
    """
    rows = list(assets)
    gap_list = list(gaps)
    out = []
    for e in catalogue:
        matches = [m for c in e["triggered_by"]
                   if (m := _match_clause(c, rows, gap_list)) is not None]
        if not matches:
            continue
        out.append(Requirement(
            id=e["id"], level=e["level"], area=e["area"], name=_flat(e["name"]),
            statement=_flat(e["statement"]),
            standards=[{k: _flat(str(s[k])) for k in ("ref", "level", "says")}
                       for s in e["standards"]],
            selection_test=_flat(e["selection_test"]),
            acceptance_test=_flat(e["acceptance_test"]),
            matches=matches))
    return out


def _flat(text: str) -> str:
    return " ".join(str(text).split())


# ── evidence ──

def _plural(n: int, unit: str) -> str:
    return f"{n} {unit}" if n == 1 else f"{n} {unit}s"


def _named(items: Sequence[str]) -> str:
    shown = list(items[:NAMED_LIMIT])
    rest = len(items) - len(shown)
    return ", ".join(shown) + (f" and {rest} more" if rest else "")


def evidence_sentence(matches: Sequence[Match]) -> str:
    """"Triggered by ..." with each clause's count, finding and the hosts behind it."""
    parts = []
    for m in matches:
        part = f"{_plural(m.count, m.unit)} where {m.finding}"
        if m.hosts:
            part += f" ({_named(m.hosts)})"
        elif m.examples:
            part += " (for example: " + "; ".join(f'"{x}"' for x in m.examples[:2]) + ")"
        parts.append(part)
    return "Triggered by " + "; ".join(parts) + "."


# ── observation input ──

def load_observation(obs_dir: str | Path) -> tuple[list[CryptoAsset], list[coverage.Gap]]:
    """The assets and coverage gaps of a stored observation.

    Gaps come from ``summary.json``'s ``gaps`` list and its parallel
    ``gap_kinds`` list; a gap with no kind reads as unclassified. A missing
    summary means no gaps.
    """
    d = Path(obs_dir)
    assets = load_assets(json.loads((d / "assets.json").read_text()))
    summary_path = d / "summary.json"
    summary = json.loads(summary_path.read_text()) if summary_path.exists() else {}
    texts = list(summary.get("gaps") or [])
    kinds = list(summary.get("gap_kinds") or [])
    gaps = [coverage.Gap(kinds[i] if i < len(kinds) and kinds[i] in coverage.KINDS
                         else coverage.UNCLASSIFIED, str(t))
            for i, t in enumerate(texts)]
    return assets, gaps


# ── output ──

def to_markdown(reqs: Sequence[Requirement], title: str = "Requirements derived from the findings",
                source: str = "") -> str:
    """The register as markdown: one ``### ID · LEVEL — name`` block per requirement."""
    lines = [f"# {title}", ""]
    if source:
        lines += [f"Derived from the observation `{source}`.", ""]
    counts = {lvl: sum(1 for r in reqs if r.level == lvl) for lvl in LEVELS}
    lines += [f"{len(reqs)} requirements were triggered: "
              + ", ".join(f"{counts[lvl]} {lvl}" for lvl in LEVELS) + ". "
              "Levels read as MUST (mandatory), SHOULD (recommended at moderate risk "
              "and above) and COULD (recommended at high risk and above).", ""]
    if not reqs:
        lines += ["No finding in this observation triggers a requirement in the catalogue.", ""]
    for r in reqs:
        lines += [f"### {r.id} · {r.level} — {r.name}", "",
                  f"- **Area:** {r.area}",
                  f"- **Statement:** {r.statement}",
                  "- **Standards:**",
                  *(f"  - {s['ref']} ({s['level']}): {s['says']}" for s in r.standards),
                  f"- **Test at selection:** {r.selection_test}",
                  f"- **Test at acceptance:** {r.acceptance_test}",
                  f"- **Evidence:** {r.evidence}",
                  ""]
    return "\n".join(lines)


REGISTER_COLUMNS = ("ID", "Level", "Area", "Requirement", "Statement", "Standards",
                    "Test at selection", "Test at acceptance", "Evidence", "Findings")
_WIDTHS = (9, 9, 18, 34, 60, 60, 45, 45, 60, 10, 22, 40, 26)


def write_xlsx(reqs: Sequence[Requirement], path: str | Path) -> None:
    """The register as a workbook: one row per requirement, supplier columns left empty."""
    # openpyxl ships no type information; its stub package is not a dependency.
    from openpyxl import Workbook  # type: ignore[import-untyped]
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # type: ignore[import-untyped]
    from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
    from openpyxl.worksheet.datavalidation import DataValidation  # type: ignore[import-untyped]

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Requirements"
    headers = REGISTER_COLUMNS + RESPONSE_COLUMNS
    ws.append(list(headers))
    for r in reqs:
        ws.append([r.id, r.level, r.area, r.name, r.statement,
                   "\n".join(f"{s['ref']} ({s['level']}): {s['says']}" for s in r.standards),
                   r.selection_test, r.acceptance_test, r.evidence, r.findings_count,
                   "", "", ""])

    thin = Side(style="thin", color="BFBFBF")
    head_fill = PatternFill("solid", fgColor="D9D9D9")
    answer_fill = PatternFill("solid", fgColor="F2F2F2")
    first_answer = len(REGISTER_COLUMNS) + 1
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if cell.row == 1:
                cell.font = Font(bold=True)
                cell.fill = head_fill
            elif cell.column >= first_answer:
                cell.fill = answer_fill
    for i, width in enumerate(_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "B2"
    if reqs:
        ws.auto_filter.ref = ws.dimensions
        dv = DataValidation(type="list", formula1=f'"{RESPONSE_CHOICES}"', allow_blank=True)
        col = get_column_letter(first_answer)
        dv.add(f"{col}2:{col}{len(reqs) + 1}")
        ws.add_data_validation(dv)
    wb.save(str(path))


def to_json(reqs: Sequence[Requirement], source: str = "") -> dict:
    """The register as plain data."""
    return {"source": source, "levels": list(LEVELS),
            "requirements": [r.to_dict() for r in reqs]}


def write_outputs(reqs: Sequence[Requirement], out_dir: str | Path,
                  source: str = "", stem: str = "requirements") -> dict[str, Path]:
    """Write ``<stem>.md``, ``<stem>.xlsx`` and ``<stem>.json`` into ``out_dir``.

    Returns the three paths keyed by format.
    """
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    paths = {"md": out / f"{stem}.md", "xlsx": out / f"{stem}.xlsx",
             "json": out / f"{stem}.json"}
    paths["md"].write_text(to_markdown(reqs, source=source))
    write_xlsx(reqs, paths["xlsx"])
    paths["json"].write_text(json.dumps(to_json(reqs, source), indent=2) + "\n")
    return paths


def main(argv: list[str] | None = None) -> int:
    """``python -m takestock.requirements <observation_dir> <out_dir> [catalogue]``.

    Exits as ``takestock requirements`` does: 1 when at least one requirement
    is derived, 0 when none is, 2 for wrong arguments or an unreadable input.
    """
    return guarded("takestock.requirements", _main, sys.argv[1:] if argv is None else argv)


def _main(args: list[str]) -> int:
    if len(args) not in (2, 3):
        print("usage: python -m takestock.requirements <observation_dir> <out_dir> "
              "[catalogue.yaml]", file=sys.stderr)
        return 2
    obs, out = Path(args[0]), Path(args[1])
    assets, gaps = load_observation(obs)
    reqs = derive(assets, load_catalogue(args[2] if len(args) == 3 else None), gaps)
    paths = write_outputs(reqs, out, source=obs.name)
    print(f"{len(reqs)} requirements triggered; wrote "
          + ", ".join(str(p) for p in paths.values()))
    return 1 if reqs else 0


if __name__ == "__main__":
    sys.exit(main())
