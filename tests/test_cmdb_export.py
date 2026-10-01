"""Writing the register: one JSON document, CSV per class, a workbook, the key ledger."""

from __future__ import annotations

import csv
import json

import cmdbfake
import pytest
from openpyxl import load_workbook

from takestock import cmdb, cmdb_export


@pytest.fixture()
def written(tmp_path):
    obs = cmdbfake.write_observation(tmp_path / "observations")
    register = cmdb.build_register(obs)
    out = tmp_path / "out"
    return register, out, cmdb.write_outputs(register, out)


def test_every_view_is_written(written):
    _register, out, paths = written
    names = {p.relative_to(out).as_posix() for p in paths}
    assert {"cmdb.json", "cmdb.xlsx", "keys.jsonl", "csv/server.csv",
            "csv/relationships.csv", "csv/contract.csv"} <= names
    assert len([n for n in names if n.startswith("csv/")]) == 26


def test_the_json_carries_the_schema_beside_the_data(written):
    register, out, _ = written
    doc = json.loads((out / "cmdb.json").read_text())
    assert list(doc)[:3] == ["register", "how_to_read", "schema"]
    assert len(doc["items"]) == len(register["items"])
    assert "key_ledger" not in doc
    attrs = {a["name"] for a in doc["schema"]["attributes"]}
    for item in doc["items"]:
        assert set(item["attributes"]) <= attrs


def test_compaction_never_drops_a_real_false(written):
    _register, out, _ = written
    doc = json.loads((out / "cmdb.json").read_text())
    assert cmdb_export._compact({"value": False, "contested": False, "note": ""}) == {"value": False}
    names = [i for i in doc["items"] if i["class"] == "DNS Name"]
    assert names and all("value" in i["attributes"]["internet_facing"]
                         or "absent" in i["attributes"]["internet_facing"] for i in names)


def test_a_class_csv_has_the_class_columns_and_no_blank_attribute(written):
    _register, out, _ = written
    with (out / "csv" / "server.csv").open(newline="") as fh:
        rows = list(csv.reader(fh))
    header = rows[0]
    attributes = [a for a in cmdb.load_schema().class_attributes("Server") if a != "lifecycle_state"]
    assert header == cmdb_export._ITEM_FIXED + attributes
    assert len(rows) == 2
    server = dict(zip(header, rows[1], strict=True))
    assert server["name"] == "app-01.example.test"
    assert server["os"] == "Linux 6" or server["os"] == "Linux"
    assert "os" in server["contested"]
    assert server["mac_address"] == "[not collected]"
    assert all(server[a] != "" for a in attributes)
    with (out / "csv" / "contract.csv").open(newline="") as fh:
        assert len(list(csv.reader(fh))) == 1


def test_the_workbook_holds_a_sheet_per_populated_class(written):
    register, out, _ = written
    wb = load_workbook(out / "cmdb.xlsx")
    populated = {c["class"] for c in register["census"] if c["items"]}
    assert populated <= set(wb.sheetnames)
    assert "Contract" not in wb.sheetnames
    for sheet in ("Read me", "Relationships", "Not in here", "Absence counted",
                  "Disagreements", "Populations", "Facts", "Data dictionary"):
        assert sheet in wb.sheetnames
    assert wb["Facts"].max_row == len(register["facts"]) + 1
    assert wb["Server"]["A2"].value.startswith("CI-")


def test_the_workbook_styling_is_neutral(written):
    _register, out, _ = written
    wb = load_workbook(out / "cmdb.xlsx")
    allowed = {None, "FF000000", "00000000", f"00{cmdb_export._ABSENT_INK}",
               f"00{cmdb_export._ACCENT}", f"FF{cmdb_export._ABSENT_INK}", f"FF{cmdb_export._ACCENT}"}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                colour = cell.font.color.rgb if cell.font and cell.font.color else None
                assert colour in allowed or not isinstance(colour, str), (ws.title, colour)
                assert cell.font.name in (None, "Calibri")


def test_a_later_build_keeps_the_keys(written, tmp_path):
    register, out, _ = written
    obs = tmp_path / "observations" / cmdbfake.STAMP
    again = cmdb.build_register(obs, keys=out / "keys.jsonl")
    assert ({i["key"]: i["name"] for i in again["items"]}
            == {i["key"]: i["name"] for i in register["items"]})


def test_the_module_runs_from_a_shell(tmp_path, capsys):
    obs = cmdbfake.write_observation(tmp_path)
    assert cmdb.main([str(obs)]) == 0
    assert (obs / "cmdb" / "cmdb.json").is_file()
    assert "items" in capsys.readouterr().out
    assert cmdb.main([str(tmp_path / "missing")]) == 2
