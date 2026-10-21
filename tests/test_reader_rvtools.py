"""The RVTools reader in takestock's readers library.

The workbooks are built here with openpyxl, in RVTools' sheet and column
names (``vInfo`` and ``vHost``, with columns RVTools documents) and its
default file name with a site label in front. Every machine and host name is
synthetic. Two sites; a template; a Site Recovery Manager placeholder named
like a machine at the other site; machines with EFI firmware and secure boot
on and off, legacy BIOS and no firmware recorded; and a blank row.
"""

from __future__ import annotations

import shutil
from datetime import UTC, datetime
from pathlib import Path

import openpyxl  # type: ignore[import-untyped]
import pytest

from takestock import config, coverage, plugin_kinds, roles
from takestock.model import CryptoAsset
from takestock.plugins import Context, StepResult
from takestock.policy import Policy
from takestock.readers import library, rvtools

LISTS = Path(__file__).resolve().parent / "fixtures" / "vsphere_lists"
VINFO = ["VM", "Powerstate", "Template", "SRM Placeholder", "Firmware", "EFI Secure boot",
         "OS according to the configuration file", "OS according to the VMware Tools"]
GUEST_2022 = "Microsoft Windows Server 2022 (64-bit)"


def _workbook(path: Path, vinfo: list[list], vhost: list[list]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "vInfo"
    for row in [VINFO, *vinfo]:
        ws.append(row)
    hosts = wb.create_sheet("vHost")
    for row in [["Host", "Datacenter", "Cluster", "ESX Version"], *vhost]:
        hosts.append(row)
    wb.save(path)
    return path


@pytest.fixture()
def exports(tmp_path) -> Path:
    folder = tmp_path / "rvtools"
    folder.mkdir()
    _workbook(folder / "SiteA_RVTools_export_all_2030-01-10_09.15.00.xlsx", [
        ["example-app-01", "poweredOn", False, False, "efi", True, GUEST_2022, GUEST_2022],
        ["example-app-02", "poweredOn", False, False, "efi", False, GUEST_2022, GUEST_2022],
        ["example-db-01", "poweredOn", False, False, "bios", None,
         "Red Hat Enterprise Linux 9 (64-bit)", ""],
        ["example-web-01", "poweredOff", False, False, "bios", None, "Ubuntu Linux (64-bit)", ""],
        ["example-tmpl-01", "poweredOff", True, False, "efi", True, GUEST_2022, ""],
        [None, None, None, None, None, None, None, None],
    ], [["esx-a1.example.test", "DC-A", "Cluster-A", "VMware ESXi 8.0.9 build-0000001"],
        ["esx-a2.example.test", "DC-A", "Cluster-A", "VMware ESXi 8.0.9 build-0000001"]])
    _workbook(folder / "SiteB_RVTools_export_all_2030-01-11_10.00.00.xlsx", [
        ["example-app-03", "poweredOn", False, False, "efi", True,
         "Microsoft Windows Server 2019 (64-bit)", ""],
        ["example-app-01", "poweredOff", False, True, "efi", False, GUEST_2022, ""],
        ["example-odd-01", "poweredOn", False, False, "", None, "", ""],
    ], [["esx-b1.example.test", "DC-B", "Cluster-B", "VMware ESXi 7.0.9 build-0000002"]])
    (folder / "notes.txt").write_text("not a workbook")
    return folder


def _ctx(path: Path | None, settings: dict | None = None) -> Context:
    return Context(plugin=rvtools.PLUGIN, target={}, settings=settings or {}, input=path,
                   ref_time=datetime(2030, 2, 1, tzinfo=UTC), raw_dir=Path("."),
                   policy=Policy(config.load_baseline()), baseline={}, dictionary={},
                   port_table={}, collected_on="2030-02-01")


def _by_name(rows: list[CryptoAsset]) -> dict[str, CryptoAsset]:
    return {a.name: a for a in rows}


def test_every_row_alone(exports):
    out = rvtools.read(_ctx(exports))
    assert out.gaps == []
    assert [a.name for a in out.records] == [
        "VM firmware EFI, secure boot on", "VM firmware EFI, secure boot off",
        "VM firmware BIOS (legacy)", "VM firmware (not reported)",
        "VMware ESXi 8.0.9 build-0000001", "VMware ESXi 7.0.9 build-0000002",
        "Guest OS Microsoft Windows Server 2022 (64-bit)",
        "Guest OS Red Hat Enterprise Linux 9 (64-bit)", "Guest OS Ubuntu Linux (64-bit)",
        "Guest OS Microsoft Windows Server 2019 (64-bit)", "Guest OS (not recorded)",
        "Virtual estate scope"]
    for a in out.records:
        assert (a.source, a.verdict, a.data_state, a.protocol, a.collected_at) == (
            "rvtools", "n/a", "platform-integrity", "firmware", "2030-01-10")
        assert a.source_name == "RVTools workbooks SiteA, SiteB, dated 2030-01-10 and 2030-01-11"
        assert plugin_kinds.check_asset("general", a) == []


def test_firmware_rows_count_machines_powered_on(exports):
    rows = _by_name(rvtools.read(_ctx(exports)).records)
    on = rows["VM firmware EFI, secure boot on"]
    assert (on.flags, on.hosts_count) == (["vm-firmware-census:2"], 2)
    assert on.extra == {"firmware": "efi", "secure_boot": "on", "powered_on_total": "5",
                        "sites": "SiteA, SiteB"}
    assert rows["VM firmware EFI, secure boot off"].flags == ["vm-firmware-census:1",
                                                              "secure-boot-off:1"]
    bios = rows["VM firmware BIOS (legacy)"]
    assert bios.flags == ["vm-firmware-census:1", "legacy-bios:1"]
    assert bios.extra["secure_boot"] == "n/a"
    assert on.role == roles.virtual_machine_boot_population(2)


def test_the_scope_row_reconciles_rows_to_machines(exports):
    scope = _by_name(rvtools.read(_ctx(exports)).records)["Virtual estate scope"]
    assert scope.flags == ["scope-rvtools-no-encryption-data"]
    assert scope.hosts_count == 6
    assert scope.extra == {"vinfo_rows": "8", "templates": "1", "srm_placeholders": "1",
                           "real_vms": "6", "distinct_vm_names": "6", "powered_on": "5",
                           "windows": "3", "non_windows": "3", "esxi_hosts": "3",
                           "export_dates": "2030-01-10 and 2030-01-11",
                           "sites": "SiteA, SiteB"}
    assert "no column for VM encryption" in scope.notes


def test_esxi_and_guest_rows(exports):
    rows = _by_name(rvtools.read(_ctx(exports)).records)
    esxi = rows["VMware ESXi 8.0.9 build-0000001"]
    assert esxi.extra == {"esx_version": "VMware ESXi 8.0.9 build-0000001",
                          "hosts": "esx-a1.example.test esx-a2.example.test", "sites": "SiteA"}
    assert (esxi.software, esxi.flags, esxi.hosts_count) == (
        "VMware ESXi 8.0.9 build-0000001", ["esxi-version-census:2"], 2)
    guest = rows["Guest OS Microsoft Windows Server 2022 (64-bit)"]
    assert guest.extra == {"guest_os": GUEST_2022, "vms": "2", "powered_on": "2"}
    assert guest.flags == ["guest-os-census:2"]


def test_with_list_exports_only_firmware_is_written(exports, tmp_path):
    lists = tmp_path / "lists"
    shutil.copytree(LISTS, lists)
    rows = rvtools.read(_ctx(exports, {"vsphere_lists": str(lists)})).records
    assert [a.name for a in rows] == [
        "VM firmware EFI, secure boot on", "VM firmware EFI, secure boot off",
        "VM firmware BIOS (legacy)", "VM firmware (not reported)", "Virtual estate scope"]
    assert rows[-1].extra == {"superseded_by": "vsphere_lists", "powered_on": "5",
                              "export_dates": "2030-01-10 and 2030-01-11",
                              "sites": "SiteA, SiteB"}


def test_one_workbook_as_the_input(exports):
    one = exports / "SiteB_RVTools_export_all_2030-01-11_10.00.00.xlsx"
    scope = _by_name(rvtools.read(_ctx(one)).records)["Virtual estate scope"]
    assert scope.extra["sites"] == "SiteB" and scope.extra["real_vms"] == "2"


def test_site_and_date_from_the_file_name(tmp_path):
    assert rvtools.site_of(Path("SiteA_RVTools_export_all_2030-01-10_09.15.00.xlsx")) == "SiteA"
    assert rvtools.site_of(Path("inventory.xlsx")) == "inventory"
    assert rvtools.date_of(Path("RVTools_export_all_2030-01-10_09.15.00.xlsx")) == "2030-01-10"


def test_no_input_writes_nothing():
    assert rvtools.read(_ctx(None)) == StepResult()


def test_workbooks_that_cannot_be_read_are_a_gap(tmp_path):
    folder = tmp_path / "rvtools"
    folder.mkdir()
    out = rvtools.read(_ctx(folder))
    [gap] = out.gaps
    assert "no .xlsx workbook" in str(gap)
    (folder / "SiteC_RVTools_export_all_2030-01-12_08.00.00.xlsx").write_bytes(b"not a zip")
    wb = openpyxl.Workbook()
    wb.active.title = "Summary"
    wb.save(folder / "SiteD_other.xlsx")
    out = rvtools.read(_ctx(folder))
    assert out.records == []
    [gap] = out.gaps
    assert "SiteC_RVTools" in str(gap) and "no vInfo rows" in str(gap)
    assert coverage.kind_of(gap) == coverage.NOT_COLLECTED
    assert coverage.nature_of(gap) == coverage.NO_DATA


def test_the_declaration():
    p = rvtools.PLUGIN
    assert (p.name, p.kind, p.input, p.gap_prefix) == ("rvtools", "general", "", False)
    assert library()["rvtools"] is p
