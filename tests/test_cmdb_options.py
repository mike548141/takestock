"""Register options: a schema overlay, export options and the Jira Assets import."""

from __future__ import annotations

import json
from pathlib import Path

import cmdbfake
import pytest
from openpyxl import load_workbook

from takestock import cmdb, cmdb_export, htmlpage

OVERLAY = """\
classes:
  Server:
    answers: Which servers does Example Widgets run?
attributes:
  serial_number:
    note: As printed on the asset label Example Widgets fixes to each machine.
relationships:
  installed-on:
    note: Drawn from the software inventory.
families:
  Hardware:
    note: Machines Example Widgets owns or leases.
vocabularies:
  absence_reason:
    not-collected:
      definition: Nobody at Example Widgets supplied the input.
governance:
  accountable_owner: Head of Operations, Example Widgets
"""


def _overlay(tmp_path: Path, text: str = OVERLAY) -> Path:
    path = tmp_path / "example-widgets.cmdb-schema-overlay.yaml"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text)
    return path


def _schema_entry(register: dict, section: str, name: str) -> dict:
    return next(e for e in register["schema"][section] if e["name"] == name)


# ── The schema overlay ─────────────────────────────────────────────────────


def test_an_overlay_rewords_the_schema_in_the_register(tmp_path):
    obs = cmdbfake.write_observation(tmp_path / "observations")
    register = cmdb.build_register(obs, schema_overlay=_overlay(tmp_path))
    assert _schema_entry(register, "classes", "Server")["answers"] == (
        "Which servers does Example Widgets run?")
    assert _schema_entry(register, "attributes", "serial_number")["note"] == (
        "As printed on the asset label Example Widgets fixes to each machine.")
    assert _schema_entry(register, "relationships", "installed-on")["note"] == (
        "Drawn from the software inventory.")
    assert register["schema"]["families"]["Hardware"]["note"] == (
        "Machines Example Widgets owns or leases.")
    assert (register["schema"]["vocabularies"]["absence_reason"]["not-collected"]["definition"]
            == "Nobody at Example Widgets supplied the input.")
    assert register["register"]["governance"]["accountable_owner"] == (
        "Head of Operations, Example Widgets")
    assert register["register"]["schema_overlay"] == "example-widgets.cmdb-schema-overlay.yaml"
    # The model itself is the bundled one: same classes, relationships and counts.
    plain = cmdb.build_register(obs)
    assert ([c["name"] for c in register["schema"]["classes"]]
            == [c["name"] for c in plain["schema"]["classes"]])
    assert register["register"]["counts"] == plain["register"]["counts"]
    assert plain["register"]["schema_overlay"] == ""


def test_the_reworded_note_reaches_the_workbook(tmp_path):
    obs = cmdbfake.write_observation(tmp_path / "observations")
    register = cmdb.build_register(obs, schema_overlay=_overlay(tmp_path))
    cmdb.write_outputs(register, tmp_path / "out")
    sheet = load_workbook(tmp_path / "out" / "cmdb.xlsx")["Data dictionary"]
    notes = {row[0].value: row[4].value for row in sheet.iter_rows(min_row=2)}
    assert notes["serial_number"] == (
        "As printed on the asset label Example Widgets fixes to each machine.")


def test_an_overlay_leaves_the_bundled_schema_alone(tmp_path):
    before = cmdb.load_schema().attributes["serial_number"].get("note")
    cmdb.load_schema(_overlay(tmp_path))
    assert cmdb.load_schema().attributes["serial_number"].get("note") == before


@pytest.mark.parametrize(("text", "says"), [
    ("classes:\n  Spaceship:\n    answers: What flies?\n", "classes.Spaceship is not declared"),
    ("relationships:\n  orbits:\n    note: Moons.\n", "relationships.orbits is not declared"),
    ("attributes:\n  colour:\n    note: Paint.\n", "attributes.colour is not declared"),
    ("families:\n  Vehicles:\n    note: Cars.\n", "families.Vehicles is not declared"),
    ("classes:\n  Server:\n    attributes: [os]\n", "classes.Server.attributes is not a wording"),
    ("classes:\n  Server:\n    family: Service\n", "classes.Server.family is not a wording"),
    ("relationships:\n  runs-on:\n    built: true\n", "relationships.runs-on.built is not a"),
    ("relationships:\n  runs-on:\n    to: Site\n", "relationships.runs-on.to is not a"),
    ("attributes:\n  os:\n    expected_strength: declared\n", "attributes.os.expected_strength"),
    ("attributes:\n  os:\n    note: [a, list]\n", "attributes.os.note must be text"),
    ("vocabularies:\n  absence_reason:\n    lost-it:\n      definition: Gone.\n",
     "vocabularies.absence_reason.lost-it is not a declared code"),
    ("vocabularies:\n  colours:\n    red:\n      definition: Red.\n",
     "vocabularies.colours is not declared"),
    ("governance:\n  treasurer: Someone\n", "governance.treasurer is not a declared"),
    ("schema_version: '9.9.9'\n", "schema_version cannot be overlaid"),
    ("authorities:\n  example-standard: {title: A standard}\n", "authorities cannot be overlaid"),
])
def test_an_overlay_that_changes_the_model_is_refused(tmp_path, text, says):
    with pytest.raises(cmdb.OverlayError, match="refused") as caught:
        cmdb.load_schema(_overlay(tmp_path, text))
    assert says in str(caught.value)


def test_a_refusal_lists_every_problem_at_once(tmp_path):
    text = ("classes:\n  Spaceship:\n    note: x\n"
            "relationships:\n  orbits:\n    note: y\n"
            "attributes:\n  os:\n    note: Fine on its own.\n")
    with pytest.raises(cmdb.OverlayError) as caught:
        cmdb.load_schema(_overlay(tmp_path, text))
    assert "Spaceship" in str(caught.value) and "orbits" in str(caught.value)


def test_a_build_with_a_refused_overlay_writes_nothing(tmp_path):
    obs = cmdbfake.write_observation(tmp_path / "observations")
    bad = _overlay(tmp_path, "classes:\n  Spaceship:\n    note: x\n")
    with pytest.raises(cmdb.OverlayError):
        cmdb.build_register(obs, schema_overlay=bad)
    with pytest.raises(cmdb.OverlayError, match="no schema overlay"):
        cmdb.load_schema(tmp_path / "missing.yaml")


# ── Options from a target profile ──────────────────────────────────────────


def _profile(tmp_path: Path, block: str, extra: str = "") -> dict:
    path = tmp_path / "targets" / "example-widgets.yaml"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("name: example-widgets\norganisation:\n  name: Example Widgets Ltd\n"
                    + extra + block)
    from takestock import config
    return config.load_target(str(path))


def test_a_profile_names_the_overlay_title_notes_and_jira_import(tmp_path):
    _overlay(tmp_path / "targets")
    (tmp_path / "targets" / "reading.yaml").write_text(
        "- heading: Start with the servers.\n  body: The Server sheet lists them.\n")
    target = _profile(tmp_path, "cmdb:\n  title: Example Widgets register\n"
                                "  schema_overlay: example-widgets.cmdb-schema-overlay.yaml\n"
                                "  how_to_read: reading.yaml\n  jira_assets: true\n")
    options = cmdb.profile_options(target)
    assert options.title == "Example Widgets register"
    assert options.organisation == "Example Widgets Ltd"
    assert options.schema_overlay == tmp_path / "targets" / "example-widgets.cmdb-schema-overlay.yaml"
    assert options.how_to_read == ({"heading": "Start with the servers.",
                                    "body": "The Server sheet lists them."},)
    assert options.jira_assets is True
    assert options.theme is None  # no branding block, so the workbook stays plain


def test_a_profile_with_branding_styles_the_workbook(tmp_path):
    target = _profile(tmp_path, "", "branding:\n  font: \"Example Sans, sans-serif\"\n"
                                    "  palette:\n    accent: \"#7a1f4d\"\n")
    assert isinstance(cmdb.profile_options(target).theme, htmlpage.Theme)


@pytest.mark.parametrize(("block", "says"), [
    ("cmdb:\n  colour: red\n", "does not take colour"),
    ("cmdb:\n  jira_assets: yes please\n", "true or false"),
    ("cmdb:\n  how_to_read: [{heading: Only a heading}]\n", "a heading and a body"),
    ("cmdb: [title]\n", "must be a mapping"),
])
def test_a_malformed_cmdb_block_is_refused(tmp_path, block, says):
    with pytest.raises(ValueError, match=says):
        cmdb.profile_options(_profile(tmp_path, block))


def test_the_command_line_reads_the_profile(tmp_path):
    from takestock import cli
    _overlay(tmp_path / "targets")
    _profile(tmp_path, "cmdb:\n  title: Example Widgets register\n"
                       "  schema_overlay: example-widgets.cmdb-schema-overlay.yaml\n"
                       "  jira_assets: true\n")
    obs = cmdbfake.write_observation(tmp_path / "observations")
    out = tmp_path / "out"
    assert cli.main(["cmdb", str(obs), str(out), "--workspace", str(tmp_path),
                     "--target", "example-widgets"]) == 0
    doc = json.loads((out / "cmdb.json").read_text())
    assert doc["register"]["title"] == "Example Widgets register"
    assert doc["register"]["organisation"] == "Example Widgets Ltd"
    assert doc["register"]["schema_overlay"] == "example-widgets.cmdb-schema-overlay.yaml"
    assert (out / "jira-assets" / "schema-and-mapping.json").is_file()
    assert cmdb.main([str(obs), str(tmp_path / "plain")]) == 0
    assert not (tmp_path / "plain" / "jira-assets").exists()


# ── Export options ─────────────────────────────────────────────────────────


def test_title_organisation_and_reading_notes_reach_the_json_and_workbook(tmp_path):
    obs = cmdbfake.write_observation(tmp_path / "observations")
    notes = [("Start with the servers.", "The Server sheet lists them."),
             {"heading": "Then the certificates.", "body": "Each has an expiry date."}]
    register = cmdb.build_register(obs, title="Example Widgets register",
                                   organisation="Example Widgets Ltd", how_to_read=notes)
    out = tmp_path / "out"
    cmdb.write_outputs(register, out)
    doc = json.loads((out / "cmdb.json").read_text())
    assert doc["register"]["title"] == "Example Widgets register"
    assert doc["register"]["organisation"] == "Example Widgets Ltd"
    assert [n["heading"] for n in doc["how_to_read"]] == ["Start with the servers.",
                                                          "Then the certificates."]
    readme = [c.value for c in load_workbook(out / "cmdb.xlsx")["Read me"]["A"] if c.value]
    assert readme[:2] == ["Example Widgets register", "Example Widgets Ltd"]
    assert "Then the certificates." in readme and "Each has an expiry date." in readme
    assert cmdb.HOW_TO_READ[0][0] not in readme


def test_without_options_the_built_in_notes_are_kept(tmp_path):
    register = cmdb.build_register(cmdbfake.write_observation(tmp_path))
    assert [n["heading"] for n in register["how_to_read"]] == [h for h, _ in cmdb.HOW_TO_READ]
    assert register["register"]["organisation"] == ""


def test_a_theme_styles_the_workbook(tmp_path):
    theme = htmlpage.theme_from_target({"branding": {
        "font": "'Example Sans', sans-serif",
        "palette": {"accent": "#7a1f4d", "surface": "#eef2f7", "text": "#1b2a3f",
                    "muted": "#4d5a6b"}}})
    register = cmdb.build_register(cmdbfake.write_observation(tmp_path),
                                   title="Example Widgets register")
    cmdb.write_outputs(register, tmp_path / "out", theme=theme)
    wb = load_workbook(tmp_path / "out" / "cmdb.xlsx")
    title = wb["Read me"]["A1"]
    assert title.font.color.rgb.endswith("7A1F4D") and title.font.name == "Example Sans"
    header = wb["Relationships"]["A1"]
    assert header.fill.fgColor.rgb.endswith("EEF2F7")
    assert header.font.color.rgb.endswith("1B2A3F")
    assert wb["Facts"]["A2"].font.name == "Example Sans"
    absent = next(c for row in wb["Server"].iter_rows(min_row=2) for c in row
                  if isinstance(c.value, str) and c.value.startswith("["))
    assert absent.font.color.rgb.endswith("4D5A6B") and absent.font.italic


def test_a_generic_font_family_names_no_workbook_font():
    assert cmdb_export.workbook_style(htmlpage.Theme()).font == ""
    assert cmdb_export.workbook_style(None) == cmdb_export.WorkbookStyle()


# ── The Jira Assets import ─────────────────────────────────────────────────

# The fields Atlassian's 2023-10-19 "Imports Schema and Mapping definition"
# marks required, and the attribute types it allows.
_REQUIRED = {"objectSchema": {"name", "description", "objectTypes"},
             "objectType": {"name", "description"},
             "objectAttribute": {"name", "description", "type"},
             "objectTypeMapping": {"objectTypeName", "selector", "description"},
             "attributeMapping": {"attributeName", "attributeLocators"}}
_TYPES = {"text", "integer", "boolean", "double", "date", "time", "date_time", "url", "email",
          "textarea", "select", "ipaddress", "referenced_object", "status"}


@pytest.fixture()
def jira(tmp_path):
    register = cmdb.build_register(cmdbfake.write_observation(tmp_path),
                                   title="Example Widgets register",
                                   organisation="Example Widgets Ltd")
    paths = cmdb.write_outputs(register, tmp_path / "out", jira_assets=True)
    out = tmp_path / "out" / "jira-assets"
    assert {out / "schema-and-mapping.json", out / "data.json"} <= set(paths)
    return (register, json.loads((out / "schema-and-mapping.json").read_text()),
            json.loads((out / "data.json").read_text()))


def test_the_mapping_has_the_fields_atlassian_requires(jira):
    _register, mapping, chunk = jira
    assert set(mapping) == {"schema", "mapping"}
    schema = mapping["schema"]["objectSchema"]
    assert _REQUIRED["objectSchema"] <= set(schema)
    assert schema["name"] == "Example Widgets register"
    assert "Example Widgets Ltd" in schema["description"]
    for object_type in schema["objectTypes"]:
        assert _REQUIRED["objectType"] <= set(object_type)
        names = [a["name"].lower() for a in object_type["attributes"]]
        assert len(names) == len(set(names)), object_type["name"]
        for attribute in object_type["attributes"]:
            assert _REQUIRED["objectAttribute"] <= set(attribute)
            assert attribute["type"] in _TYPES
    for type_mapping in mapping["mapping"]["objectTypeMappings"]:
        assert _REQUIRED["objectTypeMapping"] <= set(type_mapping)
        for attribute_mapping in type_mapping["attributesMapping"]:
            assert _REQUIRED["attributeMapping"] <= set(attribute_mapping)
        assert type_mapping["selector"] in chunk["data"]
    assert chunk["completed"] is True


def test_every_class_is_an_object_type_keyed_on_the_register_key(jira):
    register, mapping, chunk = jira
    types = {t["name"]: t for t in mapping["schema"]["objectSchema"]["objectTypes"]}
    assert set(types) == {c["name"] for c in register["schema"]["classes"]}
    server = {a["name"]: a for a in types["Server"]["attributes"]}
    assert server["Register key"]["unique"] is True and server["Name"]["label"] is True
    (server_mapping,) = [m for m in mapping["mapping"]["objectTypeMappings"]
                         if m["objectTypeName"] == "Server"]
    keyed = [m["attributeName"] for m in server_mapping["attributesMapping"] if m["externalIdPart"]]
    assert keyed == ["Register key"]
    # A class attribute called "name" is kept apart from the object's own Name.
    dns = {a["name"] for a in types["DNS Name"]["attributes"]}
    assert {"Name", "name (attribute)"} <= dns
    counts = {c["class"]: c["items"] for c in register["census"]}
    for type_mapping in mapping["mapping"]["objectTypeMappings"]:
        records = chunk["data"][type_mapping["selector"]]
        assert len(records) == counts[type_mapping["objectTypeName"]]


def test_a_relationship_becomes_a_reference_resolved_on_the_key(jira):
    register, mapping, chunk = jira
    types = {t["name"]: t for t in mapping["schema"]["objectSchema"]["objectTypes"]}
    release = {a["name"]: a for a in types["Software Release"]["attributes"]}
    reference = release["installed-on Server"]
    assert reference["type"] == "referenced_object"
    assert reference["referenceObjectTypeExternalId"] == types["Server"]["externalId"]
    assert reference["maximumCardinality"] == -1
    (release_mapping,) = [m for m in mapping["mapping"]["objectTypeMappings"]
                          if m["objectTypeName"] == "Software Release"]
    (ref_map,) = [m for m in release_mapping["attributesMapping"]
                  if m["attributeName"] == "installed-on Server"]
    locator = ref_map["attributeLocators"][0]
    assert ref_map["objectMappingIQL"] == f'"Register key" IN (${{{locator}${{0}}}})'
    (record,) = chunk["data"][release_mapping["selector"]]
    servers = {i["key"] for i in register["items"] if i["class"] == "Server"}
    assert record[locator] and set(record[locator]) <= servers
    references = {m["attributeLocators"][0] for m in release_mapping["attributesMapping"]
                  if "objectMappingIQL" in m}
    assert locator in references and len(references) == 2  # installed-on and release-of
    assert all(isinstance(v, list) == (k in references) for k, v in record.items())
    assert all(isinstance(v, str) for k, v in record.items() if k not in references)


def test_an_absence_is_written_as_its_reason(jira):
    _register, mapping, chunk = jira
    (server_mapping,) = [m for m in mapping["mapping"]["objectTypeMappings"]
                         if m["objectTypeName"] == "Server"]
    (server,) = chunk["data"][server_mapping["selector"]]
    assert server["mac_address"] == "[not collected]"
    assert server["register_name"] == "app-01.example.test"


def test_the_sample_profile_and_its_overlay_load():
    from takestock import config
    sample = Path(__file__).resolve().parents[1] / "examples" / "example-widgets" / "targets"
    options = cmdb.profile_options(config.load_target(str(sample / "example-widgets.yaml")))
    assert options.title == "Example Widgets configuration register"
    assert options.organisation == "Example Widgets Ltd" and options.jira_assets
    schema = cmdb.load_schema(options.schema_overlay)
    assert schema.overlay == "example-widgets.cmdb-schema-overlay.yaml"
    assert "asset label" in schema.attributes["serial_number"]["note"]
    # The overlay beside the profile is never taken for a profile itself.
    assert [p.name for p in config.target_profiles(sample)] == ["example-widgets.yaml"]


# ── What the read-me calls the items, and how it writes counts ─────────────


def _readme_counts(register: dict, tmp_path: Path) -> str:
    cmdb.write_outputs(register, tmp_path / "out")
    sheet = load_workbook(tmp_path / "out" / "cmdb.xlsx")["Read me"]
    return next(str(c.value) for c in sheet["A"] if c.value and "recorded facts" in str(c.value))


def test_the_read_me_counts_configuration_items_by_default(tmp_path):
    register = cmdb.build_register(cmdbfake.write_observation(tmp_path))
    assert register["register"]["item_noun"] == cmdb.ITEM_NOUN == "configuration items"
    n = register["register"]["counts"]["configuration_items"]
    assert _readme_counts(register, tmp_path).startswith(f"{n} configuration items, ")


def test_a_caller_names_the_items(tmp_path):
    register = cmdb.build_register(cmdbfake.write_observation(tmp_path),
                                   item_noun="  IT assets ")
    assert register["register"]["item_noun"] == "IT assets"
    n = register["register"]["counts"]["configuration_items"]
    line = _readme_counts(register, tmp_path)
    assert line.startswith(f"{n} IT assets, ") and "configuration items" not in line


def test_a_profile_names_the_items(tmp_path):
    options = cmdb.profile_options({"cmdb": {"item_noun": "IT assets"}})
    assert options.item_noun == "IT assets"
    assert options.build_kwargs()["item_noun"] == "IT assets"
    assert cmdb.profile_options({}).item_noun == ""


@pytest.mark.parametrize("bad", ["", "  ", None, 3, ["IT assets"]])
def test_a_profile_s_blank_or_non_text_item_noun_is_refused(bad):
    with pytest.raises(ValueError, match="cmdb.item_noun must be text"):
        cmdb.profile_options({"cmdb": {"item_noun": bad}})


def test_the_read_me_separates_thousands(tmp_path):
    register = cmdb.build_register(cmdbfake.write_observation(tmp_path))
    register["register"]["counts"].update(configuration_items=27182818, recorded_facts=3141593,
                                          relationships=1000, contested_attributes=16180339)
    line = _readme_counts(register, tmp_path)
    assert line.startswith("27,182,818 configuration items, 3,141,593 recorded facts, "
                           "1,000 relationships. ")
    assert "16,180,339 values are contested." in line
